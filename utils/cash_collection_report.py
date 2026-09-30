"""
Формирование xlsx-отчёта по строкам CashCollectionTaskTable.

Использование:

    from cash_collection_report import build_report

    build_report(rows, "report.xlsx")

где rows — список объектов CashCollectionTaskTable (или любых объектов
с такими же атрибутами), полученных, например, через
CashCollectionService.get_completed_tasks().
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FONT_NAME = "Arial"
DATE_FORMAT = "dd.mm.yyyy hh:mm"

# (заголовок столбца, имя атрибута в строке БД)
COLUMNS: Sequence[tuple[str, str]] = (
    ("Город", "city"),
    ("Адрес", "adress"),
    ("Пункт", "point"),
    ("ID аппарата", "id_device"),
    ("Описание", "description"),
    ("Комментарий сотрудника", "comment"),
    ("Сотрудник, проводивший инкассацию", "performer_name"),
    ("Задача создана в", "created_at"),
    ("Задача выполнена в", "completed_at"),
    ("Статус", "status"),
)

# столбцы, значения которых форматируются как дата/время
DATETIME_COLUMNS = {"created_at", "completed_at"}


def _sorted_rows(rows: Iterable[Any]) -> list[Any]:
    """Сортирует строки по городу, затем по дате создания для стабильного порядка."""
    return sorted(
        rows,
        key=lambda r: (
            getattr(r, "city", "") or "",
            getattr(r, "created_at", None) or datetime.min,
        ),
    )


def _write_header(ws: Worksheet) -> None:
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, (title, _attr) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    ws.freeze_panes = "A2"


def _write_rows(ws: Worksheet, rows: Sequence[Any]) -> None:
    body_font = Font(name=FONT_NAME)
    body_align = Alignment(vertical="top", wrap_text=True)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (_title, attr) in enumerate(COLUMNS, start=1):
            value = getattr(row, attr, None)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = body_font
            cell.alignment = body_align
            if attr in DATETIME_COLUMNS and value is not None:
                cell.number_format = DATE_FORMAT


def _autosize_columns(ws: Worksheet, rows: Sequence[Any]) -> None:
    for col_idx, (title, attr) in enumerate(COLUMNS, start=1):
        max_len = len(title)
        for row in rows:
            value = getattr(row, attr, None)
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        # ограничиваем ширину, чтобы очень длинные описания не растягивали лист
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)


def build_report(rows: Iterable[Any], output_path: str) -> str:
    """
    Строит xlsx-отчёт из строк CashCollectionTaskTable и сохраняет по указанному пути.
    Данные сортируются по городу.

    Возвращает output_path.
    """

    sorted_rows = _sorted_rows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Инкассация"

    _write_header(ws)
    _write_rows(ws, sorted_rows)
    _autosize_columns(ws, sorted_rows)

    wb.save(output_path)
    return output_path

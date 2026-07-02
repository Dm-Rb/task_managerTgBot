from io import BytesIO
from datetime import datetime, timedelta

from aiogram.types import BufferedInputFile

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from telegram_bot.models.task import TaskPriority


# ==========================================
# SLA
# ==========================================

PRIORITY_LIMITS = {
    TaskPriority.FAST.value: timedelta(hours=1, minutes=15),
    TaskPriority.HIGH.value: timedelta(hours=3),
    TaskPriority.NORMAL.value: timedelta(hours=8),
    TaskPriority.DAY.value: timedelta(days=1),
    TaskPriority.THREE_DAYS.value: timedelta(days=3),
}

RED_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE"
)


async def export_completed_tasks_report(repository):

    tasks = await repository.get_completed_tasks_last_30_days()

    # ==========================================
    # SORT
    # ==========================================

    tasks.sort(
        key=lambda t: (
            (t.performer_name or "").lower(),
            t.created_at or datetime.min
        )
    )

    # ==========================================
    # WORKBOOK
    # ==========================================

    wb = Workbook()
    ws = wb.active
    ws.title = "Completed Tasks"

    headers = [
        "Исполнитель",
        "Название",
        "Описание",
        "Приоритет",
        "Создана",
        "Реакция",
        "Выполнение"
    ]

    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    # ==========================================
    # ROWS
    # ==========================================

    for task in tasks:

        performer = task.performer_name or ""

        title = (
            task.title.strip()
            if task.title and task.title.strip()
            else "Без названия"
        )

        # ======================================
        # CREATED
        # ======================================

        created_at = ""

        if task.created_at:
            created_at = task.created_at.strftime(
                "%d.%m.%Y %H:%M"
            )

        # ======================================
        # REACTION
        # ======================================

        reaction_time = "-"

        if task.accepted_at and task.created_at:

            delta = (
                task.accepted_at -
                task.created_at
            )

            seconds = int(delta.total_seconds())

            hours = seconds // 3600
            minutes = (seconds % 3600) // 60

            reaction_time = (
                f"{hours:02}:{minutes:02}"
            )

            if reaction_time == "00:00":
                reaction_time = "Мгновенно"

        # ======================================
        # EXECUTION
        # ======================================

        execution_time = "-"
        full_duration = None

        if task.created_at and task.completed_at:

            full_duration = (
                task.completed_at -
                task.created_at
            )

            seconds = int(
                full_duration.total_seconds()
            )

            hours = seconds // 3600
            minutes = (
                seconds % 3600
            ) // 60

            duration = (
                f"{hours:02}:{minutes:02}"
            )

            completed_at = (
                task.completed_at.strftime(
                    "%d.%m.%Y %H:%M"
                )
            )

            if duration == "00:00":
                duration = "Мгновенно"

            execution_time = (
                f"{completed_at} ({duration})"
            )

        # ======================================
        # APPEND
        # ======================================

        ws.append([
            performer,
            title,
            task.description or "",
            task.priority,
            created_at,
            reaction_time,
            execution_time
        ])

        row = ws.max_row

        # ======================================
        # COMMENT
        # ======================================

        ws.cell(row, 2).comment = Comment(
            f"Постановщик: {task.creator_name}",
            "Bot"
        )

        # ======================================
        # SLA
        # ======================================

        limit = PRIORITY_LIMITS.get(
            task.priority
        )

        if (
                limit
                and full_duration
                and full_duration > limit
        ):
            ws.cell(row, 7).fill = RED_FILL

    # ==========================================
    # WIDTH
    # ==========================================

    for column in ws.columns:

        max_len = 0

        letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            if cell.value:

                max_len = max(
                    max_len,
                    len(str(cell.value))
                )

        ws.column_dimensions[
            letter
        ].width = min(max_len + 5, 70)

    # ==========================================
    # MEMORY
    # ==========================================

    stream = BytesIO()

    wb.save(stream)

    stream.seek(0)

    return BufferedInputFile(
        stream.read(),
        filename=(
            f"completed_tasks_"
            f"{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        )
    )